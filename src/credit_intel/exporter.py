from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document

from .config import CreditIntelConfig
from .schemas import CompanyCreditProfile


class CreditProfileExcelExporter:
    def __init__(self, config: CreditIntelConfig) -> None:
        self.config = config

    def export_company(self, profile: CompanyCreditProfile) -> Path:
        target = self.config.exports_dir / f"{_safe_name(profile.company_name)}.xlsx"
        self._write_workbook(target, [profile])
        return target

    def export_batch(self, batch_id: str, profiles: list[CompanyCreditProfile]) -> Path:
        target = self.config.batch_exports_dir / f"{batch_id}.xlsx"
        self._write_workbook(target, profiles)
        return target

    def _write_workbook(self, path: Path, profiles: list[CompanyCreditProfile]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        underwriting_rows: list[dict[str, Any]] = []
        summary_rows: list[dict[str, Any]] = []
        financial_rows: list[dict[str, Any]] = []
        rating_rows: list[dict[str, Any]] = []
        treds_rows: list[dict[str, Any]] = []
        simulation_rows: list[dict[str, Any]] = []
        advisory_rows: list[dict[str, Any]] = []
        for profile in profiles:
            simulation_payload = _simulation_payload(profile)
            underwriting_row = {
                "company_name": profile.company_name,
                "requested_name": profile.requested_name,
                "company_id": profile.company_id,
                "industry": profile.industry,
                "listed_status": profile.listed_status,
                "turnover_crore": profile.turnover_crore,
                "revenue_crore": profile.financial_summary.revenue_crore,
                "ebitda_margin_pct": profile.financial_summary.ebitda_margin_pct,
                "pat_margin_pct": profile.financial_summary.pat_margin_pct,
                "debt_to_equity": profile.financial_summary.debt_to_equity,
                "interest_coverage": profile.financial_summary.interest_coverage,
                "working_capital_days": profile.financial_summary.working_capital_days,
                "receivables_days": profile.financial_summary.receivables_days,
                "inventory_days": profile.financial_summary.inventory_days,
                "rating_available_flag": profile.rating_insight.rating_available_flag,
                "actual_rating_agency": profile.rating_insight.agency_name,
                "actual_rating": profile.rating_insight.rating,
                "actual_rating_date": profile.rating_insight.rating_date,
                "actual_rating_action": profile.rating_insight.rating_action,
                "simulation_required_flag": profile.simulation_required_flag,
                "selected_simulation_agency": simulation_payload.get("selected_agency"),
                "simulated_rating": simulation_payload.get("simulated_rating"),
                "rating_range": simulation_payload.get("rating_range"),
                "range_confidence_score": simulation_payload.get("range_confidence_score"),
                "range_confidence_label": simulation_payload.get("confidence_label"),
                "range_usability_label": simulation_payload.get("range_usability_label"),
                "simulation_actionability": simulation_payload.get("simulation_actionability"),
                "ca_review_priority": simulation_payload.get("ca_review_priority"),
                "selected_agency_reason": simulation_payload.get("selected_agency_reason"),
                "rating_history_summary": simulation_payload.get("rating_history_summary"),
                "processing_notes": " | ".join(profile.processing_notes),
            }
            underwriting_rows.append(underwriting_row)
            summary_rows.append(
                {
                    **underwriting_row,
                    "below_threshold_flag": profile.below_threshold_flag,
                    "threshold_reason": profile.threshold_reason,
                    "treds_method_used": profile.treds_insight.method_used,
                    "estimated_treds_limit_crore": profile.treds_insight.estimated_treds_limit_crore,
                    "workbook_available": profile.listed_insight.workbook_available if profile.listed_insight else False,
                    "workbook_path": profile.listed_insight.workbook_path if profile.listed_insight else None,
                }
            )
            financial_rows.append({"company_name": profile.company_name, **profile.financial_summary.model_dump()})
            for item in profile.rating_insight.history or [profile.rating_insight.model_dump()]:
                row = item.model_dump() if hasattr(item, "model_dump") else item
                row["company_name"] = profile.company_name
                rating_rows.append(row)
            treds_rows.append(
                {
                    "company_name": profile.company_name,
                    **profile.treds_insight.model_dump(),
                    "tred_platforms_detected": ", ".join(profile.treds_insight.tred_platforms_detected),
                    "tred_raw_mentions": " | ".join(profile.treds_insight.tred_raw_mentions),
                }
            )
            simulation_rows.append(
                {
                    "company_name": profile.company_name,
                    **simulation_payload,
                }
            )
            for note in profile.advisory_insight.advisory_notes or ["No rule-based advisory note generated."]:
                advisory_rows.append({"company_name": profile.company_name, "advisory_note": note})

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame(underwriting_rows).to_excel(writer, sheet_name="Underwriting View", index=False)
            pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Company Summary", index=False)
            pd.DataFrame(financial_rows).to_excel(writer, sheet_name="Financials", index=False)
            pd.DataFrame(rating_rows).to_excel(writer, sheet_name="Rating Details", index=False)
            pd.DataFrame(treds_rows).to_excel(writer, sheet_name="TReDS Insights", index=False)
            pd.DataFrame(simulation_rows).to_excel(writer, sheet_name="Simulation", index=False)
            pd.DataFrame(advisory_rows).to_excel(writer, sheet_name="Advisory Notes", index=False)


class CreditProfileWordExporter:
    def __init__(self, config: CreditIntelConfig) -> None:
        self.config = config

    def export_company(self, profile: CompanyCreditProfile) -> Path:
        target = self.config.word_exports_dir / f"{_safe_name(profile.company_name)}__credit_rationale.docx"
        self._write_document(target, profile)
        return target

    def _write_document(self, path: Path, profile: CompanyCreditProfile) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        document.add_heading(f"Credit Intelligence Draft - {profile.company_name}", level=0)
        document.add_paragraph(
            "Decision-support draft for CA/advisory use. This is not an official credit rating."
        )

        summary_table = document.add_table(rows=0, cols=2)
        summary_table.style = "Light List Accent 1"
        for label, value in (
            ("Company", profile.company_name),
            ("Requested Name", profile.requested_name),
            ("Industry", profile.industry or "Not available"),
            ("Listed Status", profile.listed_status),
            ("Turnover (Cr)", _render_value(profile.turnover_crore)),
            ("Rating", profile.rating_insight.rating or "Not available"),
            ("Agency", profile.rating_insight.agency_name or "Not available"),
            ("TReDS Method", profile.treds_insight.method_used or "Not available"),
            ("Estimated TReDS Limit (Cr)", _render_value(profile.treds_insight.estimated_treds_limit_crore)),
            ("Simulation Required", "Yes" if profile.simulation_required_flag else "No"),
        ):
            row = summary_table.add_row().cells
            row[0].text = label
            row[1].text = value

        document.add_heading("Rationale Draft", level=1)
        if profile.rationale_sections:
            for section in profile.rationale_sections:
                document.add_heading(section.heading, level=2)
                document.add_paragraph(section.body)
        elif profile.rationale_draft:
            document.add_paragraph(profile.rationale_draft)
        else:
            document.add_paragraph("No rationale text was generated for this company.")

        document.add_heading("Advisory Notes", level=1)
        if profile.advisory_insight.advisory_notes:
            for note in profile.advisory_insight.advisory_notes:
                document.add_paragraph(note, style="List Bullet")
        else:
            document.add_paragraph("No rule-based advisory note was generated.")

        document.add_heading("Processing Notes", level=1)
        if profile.processing_notes:
            for note in profile.processing_notes:
                document.add_paragraph(note, style="List Bullet")
        else:
            document.add_paragraph("No processing note was recorded.")

        document.save(path)


def _safe_name(value: str) -> str:
    return "".join(char if char.isalnum() else "_" for char in value).strip("_") or "company"


def _render_value(value: Any) -> str:
    if value is None or value == "":
        return "Not available"
    if isinstance(value, (int, float)):
        return f"{value:,.2f}"
    return str(value)


def _simulation_payload(profile: CompanyCreditProfile) -> dict[str, Any]:
    extra_payload = profile.model_extra or {}
    raw_simulation = extra_payload.get("simulation") if isinstance(extra_payload, dict) else None
    if not isinstance(raw_simulation, dict):
        return {
            "required_flag": profile.simulation_required_flag,
            "status": "Simulation unavailable",
            "simulated_rating": None,
            "rating_range": None,
            "range_confidence_score": None,
            "confidence_label": None,
            "range_usability_label": None,
            "simulation_actionability": None,
            "ca_review_priority": None,
            "selected_agency_reason": None,
            "selected_agency": None,
            "rating_history_summary": None,
            "method_used": None,
            "notes": "",
        }
    payload = dict(raw_simulation)
    payload["notes"] = " | ".join(payload.get("notes") or [])
    return payload
