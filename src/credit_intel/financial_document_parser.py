from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..utils.pdf import extract_pdf_text
from .generic_excel_parser import GenericExcelFinancialWorkbookParser
from .probe42_parser import Probe42FinancialPdfParser
from .workbook_parser import ScreenerWorkbookParser


PROVIDER_TOKENS = {
    "probe42": ("probe42.in", "probe information services private limited"),
    "scoreme": ("scoreme",),
    "finray": ("finray",),
}


class UnsupportedFinancialProviderError(ValueError):
    pass


class PendingFinancialProviderParser:
    def __init__(self, provider_name: str) -> None:
        self.provider_name = provider_name

    def parse(self, _document_path: Path) -> dict[str, Any]:
        raise NotImplementedError(
            f"{self.provider_name} ingestion is scaffolded but not implemented yet. "
            "Once we have a representative sample or API payload, we can plug it into the same canonical schema."
        )


PARSER_REGISTRY = {
    "screener": ScreenerWorkbookParser,
    "probe42": Probe42FinancialPdfParser,
    "generic_excel": GenericExcelFinancialWorkbookParser,
    "scoreme": lambda: GenericExcelFinancialWorkbookParser("scoreme"),
    "finray": lambda: GenericExcelFinancialWorkbookParser("finray"),
}


def detect_financial_provider(document_path: Path, extracted_text: str | None = None) -> str:
    suffix = document_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        workbook = load_workbook(document_path, read_only=True, data_only=True)
        sheetnames = {name.lower() for name in workbook.sheetnames}
        if "data sheet" in sheetnames:
            return "screener"
        return "generic_excel"

    lowered = (extracted_text or "").lower()
    if suffix == ".pdf":
        if not lowered:
            lowered = extract_pdf_text(document_path).text.lower()
        for provider_name, tokens in PROVIDER_TOKENS.items():
            if any(token in lowered for token in tokens):
                return provider_name
    raise UnsupportedFinancialProviderError(f"Unable to detect financial provider for {document_path}")


def get_financial_parser(provider_name: str):
    factory = PARSER_REGISTRY.get(provider_name)
    if not factory:
        raise UnsupportedFinancialProviderError(f"Unsupported financial provider: {provider_name}")
    return factory()


def parse_financial_document(document_path: Path, provider_name: str | None = None) -> dict[str, Any]:
    document_path = Path(document_path)
    detected_provider = provider_name or detect_financial_provider(document_path)
    parser = get_financial_parser(detected_provider)
    parsed = parser.parse(document_path)
    meta = dict(parsed.get("meta") or {})
    meta.setdefault("source_provider", detected_provider)
    parsed["meta"] = meta
    return parsed
