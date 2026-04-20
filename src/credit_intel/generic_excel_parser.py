from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .schemas import FinancialSummary


SHEET_EXCLUDE_TOKENS = ("director", "charge", "charges")

HEADER_SCORE_ALIASES = {
    "company name",
    "company id",
    "comp id",
    "cin",
    "financial year",
    "latest balance sheet",
    "total revenue",
    "ebitda",
    "pat",
    "networth",
    "total borrowings",
    "current ratio",
    "receivables days",
    "inventory days",
}

COLUMN_ALIASES = {
    "company_name": {"company name"},
    "company_id": {"company id", "comp id"},
    "cin": {"cin"},
    "pan": {"pan"},
    "city": {"city"},
    "state": {"state"},
    "status": {"status"},
    "listing_status": {"listing status"},
    "products": {"products"},
    "industry": {"corpository sector", "industry", "sector"},
    "whether_exporter": {"whether exporter", "importer exporter"},
    "business_type": {"business type"},
    "paid_up_capital": {"paid up capital"},
    "financial_year": ("financial year", "fy as per db"),
    "latest_balance_sheet": ("latest balance sheet", "date of balance sheet"),
    "standalone_or_consolidated": {"consolidated standalone"},
    "credit_rated_flag": {"credit rated"},
    "latest_ratings_text": {"latest ratings"},
    "total_open_charges": {"total open charges"},
    "revenue_crore": ("total revenue", "total revenue from operations"),
    "ebitda_crore": {"ebitda"},
    "pat_crore": {"pat"},
    "networth_crore": {"networth"},
    "long_term_liabilities_crore": {"long term liabilities"},
    "total_borrowings_crore": {"total borrowings"},
    "long_term_borrowings_crore": {"long term borrowings"},
    "short_term_borrowings_crore": {"short term borrowings"},
    "cash_and_cash_equivalents_crore": {"cash and cash equivalents"},
    "current_investment_crore": {"current investment"},
    "total_tangible_assets_crore": {"total tangible assets"},
    "fixed_assets_crore": {"fixed assets"},
    "revenue_growth_pct": {"revenue growth pct"},
    "finance_cost_pct_of_sales": {"finance cost pct of sales"},
    "current_ratio": {"current ratio"},
    "receivables_days": {"receivables days"},
    "inventory_days": {"inventory days"},
    "debt_to_equity": ("total debt equity", "debt equity"),
    "debt_to_ebitda": {"debt to ebitda"},
    "interest_coverage": {"interest coverage"},
    "growth_in_fixed_assets_pct": {"growth in fixed assets pct"},
    "ebitda_margin_pct": {"ebitda pct"},
    "ebt_margin_pct": {"ebt margins pct", "ebt margin pct"},
    "pat_margin_pct": {"pat pct"},
    "emp_cost_pct_of_sales": {"emp cost pct of sales"},
    "other_exp_pct_of_sales": {"other exp pct of sales"},
    "rmc_pct": {"rmc pct"},
    "roe_pct": {"roe pct"},
    "fa_turnover": {"fa t o"},
    "working_capital_turnover": {"working capital turnover"},
    "long_term_debt_to_equity": {"long term debt equity"},
}

PERCENT_KEYS = {
    "revenue_growth_pct",
    "finance_cost_pct_of_sales",
    "ebitda_margin_pct",
    "ebt_margin_pct",
    "pat_margin_pct",
    "emp_cost_pct_of_sales",
    "other_exp_pct_of_sales",
    "rmc_pct",
    "roe_pct",
    "growth_in_fixed_assets_pct",
}

PROFIT_LOSS_FIELD_MAP = {
    "sales": "revenue_crore",
    "operating_profit": "ebitda_crore",
    "net_profit": "pat_crore",
}

BALANCE_SHEET_FIELD_MAP = {
    "equity_share_capital": "paid_up_capital",
    "borrowings": "total_borrowings_crore",
    "long_term_borrowings": "long_term_borrowings_crore",
    "short_term_borrowings": "short_term_borrowings_crore",
    "cash_and_bank": "cash_and_cash_equivalents_crore",
    "investments": "current_investment_crore",
    "fixed_assets": "fixed_assets_crore",
}

RATIO_FIELD_MAP = {
    "revenue_growth_pct": "revenue_growth_pct",
    "ebitda_margin_pct": "ebitda_margin_pct",
    "pat_margin_pct": "pat_margin_pct",
    "current_ratio": "current_ratio",
    "receivables_days": "receivables_days",
    "inventory_days": "inventory_days",
    "debt_to_equity": "debt_to_equity",
    "debt_to_ebitda": "debt_to_ebitda",
    "interest_coverage": "interest_coverage",
    "roe_pct": "roe_pct",
    "sales_to_net_fixed_assets": "fa_turnover",
    "working_capital_turnover": "working_capital_turnover",
    "long_term_debt_to_equity": "long_term_debt_to_equity",
}

FINANCIAL_PARAMETER_FIELD_MAP = {
    "total_open_charges": "total_open_charges",
}


class GenericExcelFinancialWorkbookParser:
    def __init__(self, provider_name: str = "generic_excel") -> None:
        self.provider_name = provider_name

    def parse(self, workbook_path: Path) -> dict[str, Any]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_name, header_row_index, headers = self._select_best_sheet(workbook)
        worksheet = workbook[sheet_name]
        companies: list[dict[str, Any]] = []
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_index <= header_row_index:
                continue
            row_values = list(row)
            if all(value is None or str(value).strip() == "" for value in row_values):
                continue
            header_values = _collect_header_values(headers, row_values)
            payload = self._parse_company_row(header_values)
            if payload:
                companies.append(payload)

        return {
            "meta": {
                "source_provider": self.provider_name,
                "sheet_name": sheet_name,
                "header_row_index": header_row_index,
                "company_count": len(companies),
                "document_type": "generic_excel_tabular_snapshot",
            },
            "companies": companies,
        }

    def _select_best_sheet(self, workbook) -> tuple[str, int, list[str]]:  # type: ignore[no-untyped-def]
        best_score = -1
        best_sheet_name: str | None = None
        best_header_index = 1
        best_headers: list[str] = []

        for sheet_name in workbook.sheetnames:
            normalized_sheet = sheet_name.lower()
            if any(token in normalized_sheet for token in SHEET_EXCLUDE_TOKENS):
                continue
            worksheet = workbook[sheet_name]
            candidate_rows = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                candidate_rows.append((row_index, list(row)))
                if row_index >= 8:
                    break
            for row_index, row in candidate_rows:
                headers = _resolve_headers(candidate_rows, row_index)
                score = sum(1 for header in headers if header in HEADER_SCORE_ALIASES)
                if "company name" in headers:
                    score += 5
                if "total revenue" in headers or "total revenue from operations" in headers:
                    score += 3
                if score > best_score:
                    best_score = score
                    best_sheet_name = sheet_name
                    best_header_index = row_index
                    best_headers = headers

        if best_sheet_name is None or best_score < 3:
            raise ValueError("Unable to identify a tabular financial sheet with recognizable headers.")
        return best_sheet_name, best_header_index, best_headers

    def _parse_company_row(self, header_values: dict[str, list[Any]]) -> dict[str, Any] | None:
        company_name = _pick_string(header_values, COLUMN_ALIASES["company_name"])
        if not company_name:
            return None

        meta = {
            "company_name": company_name,
            "company_id": _pick_string(header_values, COLUMN_ALIASES["company_id"]),
            "cin": _pick_string(header_values, COLUMN_ALIASES["cin"]),
            "pan": _pick_string(header_values, COLUMN_ALIASES["pan"]),
            "city": _pick_string(header_values, COLUMN_ALIASES["city"]),
            "state": _pick_string(header_values, COLUMN_ALIASES["state"]),
            "status": _pick_string(header_values, COLUMN_ALIASES["status"]),
            "listing_status": _pick_string(header_values, COLUMN_ALIASES["listing_status"]),
            "products": _pick_string(header_values, COLUMN_ALIASES["products"]),
            "industry": _pick_string(header_values, COLUMN_ALIASES["industry"]),
            "whether_exporter": _pick_bool_like(header_values, COLUMN_ALIASES["whether_exporter"]),
            "business_type": _pick_string(header_values, COLUMN_ALIASES["business_type"]),
            "financial_year": _pick_string(header_values, COLUMN_ALIASES["financial_year"]),
            "latest_balance_sheet": _pick_date_string(header_values, COLUMN_ALIASES["latest_balance_sheet"]),
            "standalone_or_consolidated": _pick_string(header_values, COLUMN_ALIASES["standalone_or_consolidated"]),
            "credit_rated_flag": _pick_bool_like(header_values, COLUMN_ALIASES["credit_rated_flag"]),
            "latest_ratings_text": _pick_string(header_values, COLUMN_ALIASES["latest_ratings_text"]),
            "total_open_charges": _pick_float(header_values, COLUMN_ALIASES["total_open_charges"]),
            "source_provider": self.provider_name,
        }

        period = _derive_period(meta["latest_balance_sheet"], meta["financial_year"])
        paid_up_capital = _pick_float(header_values, COLUMN_ALIASES["paid_up_capital"])
        networth = _pick_float(header_values, COLUMN_ALIASES["networth_crore"])
        reserves = None
        if networth is not None and paid_up_capital is not None:
            reserves = round(networth - paid_up_capital, 6)

        pbt = _derive_profit_before_tax(
            sales=_pick_float(header_values, COLUMN_ALIASES["revenue_crore"]),
            ebt_margin_pct=_pick_numeric_field(header_values, "ebt_margin_pct"),
        )
        interest = _derive_from_sales_pct(
            sales=_pick_float(header_values, COLUMN_ALIASES["revenue_crore"]),
            pct_value=_pick_numeric_field(header_values, "finance_cost_pct_of_sales"),
        )
        tax = _derive_tax(pbt, _pick_float(header_values, COLUMN_ALIASES["pat_crore"]))

        financial_summary = FinancialSummary(
            revenue_crore=_pick_float(header_values, COLUMN_ALIASES["revenue_crore"]),
            ebitda_margin_pct=_pick_numeric_field(header_values, "ebitda_margin_pct"),
            pat_margin_pct=_pick_numeric_field(header_values, "pat_margin_pct"),
            debt_to_equity=_pick_numeric_field(header_values, "debt_to_equity"),
            interest_coverage=_pick_numeric_field(header_values, "interest_coverage"),
            working_capital_days=None,
            receivables_days=_pick_numeric_field(header_values, "receivables_days"),
            inventory_days=_pick_numeric_field(header_values, "inventory_days"),
            current_price=None,
            market_cap_crore=None,
            networth_crore=networth,
            total_borrowings_crore=_pick_float(header_values, COLUMN_ALIASES["total_borrowings_crore"]),
            source=f"{self.provider_name}_snapshot",
            statement_period=period,
        ).model_dump()

        payables_days = None
        working_capital_days = _compute_working_capital_days(
            receivables_days=financial_summary.get("receivables_days"),
            inventory_days=financial_summary.get("inventory_days"),
            payables_days=payables_days,
        )
        financial_summary["working_capital_days"] = working_capital_days

        profit_loss = [
            {
                "period": period,
                **_mapped_fields(header_values, PROFIT_LOSS_FIELD_MAP),
                "profit_before_tax": pbt,
                "interest": interest,
                "tax": tax,
            }
        ]

        balance_sheet = [
            {
                "period": period,
                **_mapped_fields(header_values, BALANCE_SHEET_FIELD_MAP),
                "reserves": reserves,
                "networth": networth,
            }
        ]

        ratios = [
            {
                "period": period,
                **_mapped_fields(header_values, RATIO_FIELD_MAP, numeric_lookup=_pick_numeric_field),
            }
        ]

        financial_parameters = [
            {
                "period": period,
                **_mapped_fields(header_values, FINANCIAL_PARAMETER_FIELD_MAP),
                "credit_rated_flag": meta["credit_rated_flag"],
                "latest_ratings_text": meta["latest_ratings_text"],
                "whether_exporter": meta["whether_exporter"],
            }
        ]

        payload = {
            "meta": meta,
            "financial_summary": financial_summary,
            "profit_loss": profit_loss,
            "balance_sheet": balance_sheet,
            "cash_flow": [],
            "ratios": ratios,
            "financial_parameters": financial_parameters,
        }
        return payload


def _collect_header_values(headers: list[str], row_values: list[Any]) -> dict[str, list[Any]]:
    values_by_header: dict[str, list[Any]] = {}
    for header, value in zip(headers, row_values, strict=False):
        if not header:
            continue
        values_by_header.setdefault(header, []).append(value)
    return values_by_header


def _resolve_headers(candidate_rows: list[tuple[int, list[Any]]], target_row_index: int) -> list[str]:
    rows_by_index = {row_index: row for row_index, row in candidate_rows}
    base_row = rows_by_index[target_row_index]
    headers = [_normalize_header(cell) for cell in base_row]
    for column_index, header in enumerate(headers):
        if header:
            continue
        fallback = ""
        for previous_row_index in range(target_row_index - 1, 0, -1):
            previous_row = rows_by_index.get(previous_row_index)
            if previous_row is None or column_index >= len(previous_row):
                continue
            candidate = _normalize_header(previous_row[column_index])
            if candidate:
                fallback = candidate
                break
        headers[column_index] = fallback
    return headers


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if not text:
        return ""
    text = text.replace("%", " pct ")
    text = text.replace("&", " and ")
    text = text.replace("/", " ")
    text = text.replace("*", " ")
    text = text.replace("(", " ").replace(")", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _pick_values(header_values: dict[str, list[Any]], aliases: set[str]) -> list[Any]:
    matches: list[Any] = []
    for alias in aliases:
        matches.extend(header_values.get(alias, []))
    return matches


def _pick_string(header_values: dict[str, list[Any]], aliases: set[str]) -> str | None:
    for value in _pick_values(header_values, aliases):
        if value is None:
            continue
        text = str(value).strip()
        if text and text not in {"-", "None"}:
            return text
    return None


def _pick_float(header_values: dict[str, list[Any]], aliases: set[str]) -> float | None:
    for value in _pick_values(header_values, aliases):
        parsed = _coerce_numeric(value, is_percent=False)
        if parsed is not None:
            return parsed
    return None


def _pick_numeric_field(header_values: dict[str, list[Any]], field_name: str) -> float | None:
    aliases = COLUMN_ALIASES[field_name]
    is_percent = field_name in PERCENT_KEYS
    for value in _pick_values(header_values, aliases):
        parsed = _coerce_numeric(value, is_percent=is_percent)
        if parsed is not None:
            return parsed
    return None


def _pick_bool_like(header_values: dict[str, list[Any]], aliases: set[str]) -> bool | None:
    text = _pick_string(header_values, aliases)
    if text is None:
        return None
    lowered = text.strip().lower()
    if lowered in {"yes", "y", "true"}:
        return True
    if lowered in {"no", "n", "false"}:
        return False
    return None


def _pick_date_string(header_values: dict[str, list[Any]], aliases: set[str]) -> str | None:
    for value in _pick_values(header_values, aliases):
        parsed = _coerce_date(value)
        if parsed:
            return parsed
    return None


def _coerce_numeric(value: Any, *, is_percent: bool) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        numeric = float(value)
        if is_percent and abs(numeric) <= 1.0:
            return round(numeric * 100.0, 6)
        return round(numeric, 6)
    text = str(value).strip()
    if not text or text in {"-", "None", "nan"}:
        return None
    has_percent_symbol = "%" in text
    text = text.replace(",", "").replace("%", "").strip()
    try:
        numeric = float(text)
    except ValueError:
        return None
    if is_percent and not has_percent_symbol and abs(numeric) <= 1.0:
        numeric *= 100.0
    return round(numeric, 6)


def _coerce_date(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    if not text or text in {"-", "None"}:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _derive_period(latest_balance_sheet: str | None, financial_year: str | None) -> str | None:
    if latest_balance_sheet:
        return latest_balance_sheet
    if not financial_year:
        return None
    match = re.search(r"(\d{4})\D+(\d{4})", financial_year)
    if match:
        return f"{match.group(2)}-03-31"
    return None


def _mapped_fields(
    header_values: dict[str, list[Any]],
    field_map: dict[str, str],
    *,
    numeric_lookup=None,
) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for target_field, source_key in field_map.items():
        if numeric_lookup:
            values[target_field] = numeric_lookup(header_values, source_key)
        else:
            aliases = COLUMN_ALIASES[source_key]
            values[target_field] = _pick_float(header_values, aliases)
    return values


def _derive_profit_before_tax(sales: float | None, ebt_margin_pct: float | None) -> float | None:
    if sales is None or ebt_margin_pct is None:
        return None
    return round((sales * ebt_margin_pct) / 100.0, 6)


def _derive_from_sales_pct(sales: float | None, pct_value: float | None) -> float | None:
    if sales is None or pct_value is None:
        return None
    return round((sales * pct_value) / 100.0, 6)


def _derive_tax(pbt: float | None, pat: float | None) -> float | None:
    if pbt is None or pat is None:
        return None
    return round(pbt - pat, 6)


def _compute_working_capital_days(
    receivables_days: float | None,
    inventory_days: float | None,
    payables_days: float | None,
) -> float | None:
    if receivables_days is None and inventory_days is None:
        return None
    base = float(receivables_days or 0.0) + float(inventory_days or 0.0)
    if payables_days is not None:
        return round(base - float(payables_days), 4)
    return round(base, 4)
