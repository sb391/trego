from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Pt

from .config import CreditIntelConfig
from .ui_schemas import CompanyDashboardView


DISCLAIMER_TEXT = (
    "This report is a decision-support document for advisors and is not an official credit rating. "
    "It formats backend outputs without independently computing or validating rating conclusions."
)


class PresentationReportExporter:
    def __init__(self, config: CreditIntelConfig) -> None:
        self.config = config

    @property
    def base_dir(self) -> Path:
        return self.config.exports_dir / "ui_reports"

    @property
    def excel_dir(self) -> Path:
        return self.base_dir / "excel"

    @property
    def pdf_dir(self) -> Path:
        return self.base_dir / "pdf"

    @property
    def word_dir(self) -> Path:
        return self.base_dir / "word"

    @property
    def zip_dir(self) -> Path:
        return self.base_dir / "zip"

    def export_excel(self, company: CompanyDashboardView) -> Path:
        target = self.excel_dir / f"{_safe_name(company.company_name)}__dashboard.xlsx"
        target.parent.mkdir(parents=True, exist_ok=True)

        summary_rows = [
            {
                "company_name": company.company_name,
                "requested_name": company.requested_name,
                "company_id": company.company_id,
                "industry": company.industry,
                "listed_status": company.listed_status,
                "revenue": company.revenue,
                "turnover_crore": company.turnover_crore,
                "risk_indicator": company.risk_indicator.label,
                "rating": company.ratings.rating,
                "agency": company.ratings.agency,
                "simulation_status": company.simulation.status,
                "processed_at": company.processed_at,
            }
        ]

        financial_rows: list[dict[str, Any]] = []
        for statement_type, rows in company.financials.tables.items():
            for row in rows:
                financial_rows.append({"statement_type": statement_type, **row})
        if not financial_rows:
            financial_rows.append({"statement_type": "summary", **company.financials.summary.model_dump()})

        ratings_rows = [
            {
                "rating_date": row.rating_date,
                "agency": row.agency_name,
                "rating": row.rating,
                "outlook": row.outlook,
                "rating_action": row.rating_action,
                "source_url": row.source_url,
                "source": row.source,
            }
            for row in company.ratings.history
        ] or [
            {
                "rating_date": company.ratings.rating_date,
                "agency": company.ratings.agency,
                "rating": company.ratings.rating,
                "outlook": company.ratings.outlook,
                "rating_action": company.ratings.rating_action,
                "source_url": None,
                "source": "summary",
            }
        ]

        treds_rows = [
            {
                "presence_flag": company.tred.presence_flag,
                "status": company.tred.status,
                "platforms": ", ".join(company.tred.platforms),
                "signal_strength": company.tred.signal_strength,
                "estimated_limit_crore": company.tred.estimated_limit_crore,
                "estimation_confidence": company.tred.estimation_confidence,
                "method_used": company.tred.method_used,
                "raw_mentions": " | ".join(company.tred.raw_mentions),
            }
        ]

        simulation_rows = [
            {
                "required_flag": company.simulation.required_flag,
                "status": company.simulation.status,
                "selected_agency": company.simulation.selected_agency,
                "simulated_rating": company.simulation.simulated_rating,
                "rating_range": company.simulation.rating_range,
                "confidence_score": company.simulation.confidence_score,
                "confidence_label": company.simulation.confidence_label,
                "range_usability_label": company.simulation.range_usability_label,
                "simulation_actionability": company.simulation.simulation_actionability,
                "ca_review_priority": company.simulation.ca_review_priority,
                "selected_agency_reason": company.simulation.selected_agency_reason,
                "rating_history_summary": company.simulation.rating_history_summary,
                "method_used": company.simulation.method_used,
                "notes": " | ".join(company.simulation.notes),
            }
        ]

        with pd.ExcelWriter(target, engine="openpyxl") as writer:
            pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False)
            pd.DataFrame(financial_rows).to_excel(writer, sheet_name="Financials", index=False)
            pd.DataFrame(ratings_rows).to_excel(writer, sheet_name="Ratings", index=False)
            pd.DataFrame(treds_rows).to_excel(writer, sheet_name="TReDS", index=False)
            pd.DataFrame(simulation_rows).to_excel(writer, sheet_name="Simulation", index=False)

        return target

    def export_portfolio_excel(self, batch_id: str, companies: list[CompanyDashboardView]) -> Path:
        target = self.excel_dir / f"{_safe_name(batch_id)}__portfolio.xlsx"
        target.parent.mkdir(parents=True, exist_ok=True)

        summary_rows = [
            {
                "company_id": company.company_id,
                "company_name": company.company_name,
                "requested_name": company.requested_name,
                "industry": company.industry,
                "listed_status": company.listed_status,
                "revenue": company.revenue,
                "ebitda_margin_pct": company.financials.summary.ebitda_margin_pct,
                "debt_to_equity": company.financials.summary.debt_to_equity,
                "latest_rating": company.ratings.rating,
                "rating_agency": company.ratings.agency,
                "treds_status": company.tred.status,
                "estimated_treds_limit_crore": company.tred.estimated_limit_crore,
                "simulation_range": company.simulation.rating_range,
                "risk_indicator": company.risk_indicator.label,
                "processed_at": company.processed_at,
            }
            for company in companies
        ]

        financial_rows: list[dict[str, Any]] = []
        ratings_rows: list[dict[str, Any]] = []
        treds_rows: list[dict[str, Any]] = []
        simulation_rows: list[dict[str, Any]] = []

        for company in companies:
            financial_rows.extend(_build_portfolio_financial_rows(company))
            ratings_rows.extend(_build_portfolio_ratings_rows(company))
            treds_rows.append(
                {
                    "company_id": company.company_id,
                    "company_name": company.company_name,
                    "status": company.tred.status,
                    "presence_flag": company.tred.presence_flag,
                    "platforms": ", ".join(company.tred.platforms),
                    "signal_strength": company.tred.signal_strength,
                    "estimated_limit_crore": company.tred.estimated_limit_crore,
                    "estimation_confidence": company.tred.estimation_confidence,
                    "method_used": company.tred.method_used,
                    "raw_mentions": " | ".join(company.tred.raw_mentions),
                }
            )
            simulation_rows.append(
                {
                    "company_id": company.company_id,
                    "company_name": company.company_name,
                    "required_flag": company.simulation.required_flag,
                    "status": company.simulation.status,
                    "selected_agency": company.simulation.selected_agency,
                    "simulated_rating": company.simulation.simulated_rating,
                    "rating_range": company.simulation.rating_range,
                    "confidence_score": company.simulation.confidence_score,
                    "confidence_label": company.simulation.confidence_label,
                    "range_usability_label": company.simulation.range_usability_label,
                    "simulation_actionability": company.simulation.simulation_actionability,
                    "ca_review_priority": company.simulation.ca_review_priority,
                    "selected_agency_reason": company.simulation.selected_agency_reason,
                    "rating_history_summary": company.simulation.rating_history_summary,
                    "method_used": company.simulation.method_used,
                    "notes": " | ".join(company.simulation.notes),
                }
            )

        with pd.ExcelWriter(target, engine="openpyxl") as writer:
            pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False)
            pd.DataFrame(financial_rows).to_excel(writer, sheet_name="Financials", index=False)
            pd.DataFrame(ratings_rows).to_excel(writer, sheet_name="Ratings", index=False)
            pd.DataFrame(treds_rows).to_excel(writer, sheet_name="TReDS", index=False)
            pd.DataFrame(simulation_rows).to_excel(writer, sheet_name="Simulation", index=False)

        return target

    def export_portfolio_zip(self, batch_id: str, companies: list[CompanyDashboardView]) -> Path:
        target = self.zip_dir / f"{_safe_name(batch_id)}__portfolio_reports.zip"
        target.parent.mkdir(parents=True, exist_ok=True)

        summary_workbook = self.export_portfolio_excel(batch_id, companies)

        with zipfile.ZipFile(target, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.write(summary_workbook, arcname="portfolio_reports/summary.xlsx")
            used_folder_names: set[str] = set()
            for index, company in enumerate(companies, start=1):
                folder_name = _unique_safe_name(company.company_name, used_folder_names, fallback=f"company_{index}")
                pdf_path = self.export_pdf(company)
                word_path = self.export_word(company)
                archive.write(pdf_path, arcname=f"portfolio_reports/{folder_name}/{pdf_path.name}")
                archive.write(word_path, arcname=f"portfolio_reports/{folder_name}/{word_path.name}")

        return target

    def export_word(self, company: CompanyDashboardView) -> Path:
        target = self.word_dir / f"{_safe_name(company.company_name)}__dashboard.docx"
        target.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        document.core_properties.title = f"{company.company_name} Credit Intelligence Report"
        _set_normal_font(document)

        document.add_heading(company.company_name, level=0)
        document.add_paragraph("Credit Intelligence Dashboard Report")
        self._add_heading(document, "Executive Summary")
        for line in _executive_summary_lines(company):
            document.add_paragraph(line, style="List Bullet")

        self._add_heading(document, "Financial Summary")
        financial_table = document.add_table(rows=1, cols=2)
        financial_table.style = "Light List Accent 1"
        financial_table.rows[0].cells[0].text = "Metric"
        financial_table.rows[0].cells[1].text = "Value"
        for label, value in _financial_metric_rows(company).items():
            cells = financial_table.add_row().cells
            cells[0].text = label
            cells[1].text = value

        self._add_heading(document, "Rating History")
        if company.ratings.history:
            rating_table = document.add_table(rows=1, cols=5)
            rating_table.style = "Light List Accent 1"
            headers = ["Date", "Agency", "Rating", "Outlook", "Action"]
            for index, header in enumerate(headers):
                rating_table.rows[0].cells[index].text = header
            for entry in company.ratings.history:
                cells = rating_table.add_row().cells
                cells[0].text = entry.rating_date or "—"
                cells[1].text = entry.agency_name or "—"
                cells[2].text = entry.rating or "—"
                cells[3].text = entry.outlook or "—"
                cells[4].text = entry.rating_action or "—"
        else:
            document.add_paragraph("No rating history available in the current response.")

        self._add_heading(document, "TReDS Insights")
        for line in _treds_lines(company):
            document.add_paragraph(line, style="List Bullet")

        self._add_heading(document, "Simulation Output")
        for line in _simulation_lines(company):
            document.add_paragraph(line, style="List Bullet")

        self._add_heading(document, "Our Rationale")
        if company.advisory.rationale_sections:
            for section in company.advisory.rationale_sections:
                document.add_heading(section.heading, level=2)
                document.add_paragraph(section.body)
        elif company.advisory.rationale_draft:
            document.add_paragraph(company.advisory.rationale_draft)
        else:
            document.add_paragraph("No rationale narrative was available in the current response.")

        self._add_heading(document, "Advisory")
        for note in company.advisory.notes or ["No advisory notes available."]:
            document.add_paragraph(note, style="List Bullet")

        self._add_heading(document, "Disclaimer")
        document.add_paragraph(DISCLAIMER_TEXT)

        document.save(target)
        return target

    def export_pdf(self, company: CompanyDashboardView) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("reportlab is required for PDF exports.") from exc

        target = self.pdf_dir / f"{_safe_name(company.company_name)}__dashboard.pdf"
        target.parent.mkdir(parents=True, exist_ok=True)

        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="SectionHeading", parent=styles["Heading2"], textColor=colors.HexColor("#0f1720")))
        styles.add(ParagraphStyle(name="BodySmall", parent=styles["BodyText"], leading=14))

        story: list[Any] = [
            Paragraph(company.company_name, styles["Title"]),
            Paragraph("Credit Intelligence Dashboard Report", styles["Heading3"]),
            Spacer(1, 0.18 * inch),
            Paragraph("Executive Summary", styles["SectionHeading"]),
        ]
        for line in _executive_summary_lines(company):
            story.append(Paragraph(f"- {line}", styles["BodySmall"]))
        story.append(Spacer(1, 0.18 * inch))

        story.append(Paragraph("Financial Summary", styles["SectionHeading"]))
        financial_rows = [["Metric", "Value"], *[[label, value] for label, value in _financial_metric_rows(company).items()]]
        financial_table = Table(financial_rows, colWidths=[2.4 * inch, 3.6 * inch])
        financial_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
                ]
            )
        )
        story.extend([financial_table, Spacer(1, 0.18 * inch)])

        story.append(Paragraph("Rating History", styles["SectionHeading"]))
        rating_rows = [["Date", "Agency", "Rating", "Outlook", "Action"]]
        if company.ratings.history:
            for entry in company.ratings.history[:12]:
                rating_rows.append(
                    [
                        entry.rating_date or "—",
                        entry.agency_name or "—",
                        entry.rating or "—",
                        entry.outlook or "—",
                        entry.rating_action or "—",
                    ]
                )
        else:
            rating_rows.append(["—", "—", "No rating history", "—", "—"])
        rating_table = Table(rating_rows, colWidths=[1.0 * inch, 1.1 * inch, 1.1 * inch, 1.0 * inch, 2.3 * inch])
        rating_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ]
            )
        )
        story.extend([rating_table, Spacer(1, 0.18 * inch)])

        for heading, lines in (
            ("TReDS Insights", _treds_lines(company)),
            ("Simulation Output", _simulation_lines(company)),
            ("Advisory", company.advisory.notes or ["No advisory notes available."]),
        ):
            story.append(Paragraph(heading, styles["SectionHeading"]))
            for line in lines:
                story.append(Paragraph(f"- {line}", styles["BodySmall"]))
            story.append(Spacer(1, 0.12 * inch))

        story.append(Paragraph("Our Rationale", styles["SectionHeading"]))
        if company.advisory.rationale_sections:
            for section in company.advisory.rationale_sections:
                story.append(Paragraph(section.heading, styles["Heading3"]))
                story.append(Paragraph(section.body, styles["BodySmall"]))
        elif company.advisory.rationale_draft:
            story.append(Paragraph(company.advisory.rationale_draft, styles["BodySmall"]))
        else:
            story.append(Paragraph("No rationale narrative was available in the current response.", styles["BodySmall"]))

        story.extend(
            [
                Spacer(1, 0.18 * inch),
                Paragraph("Disclaimer", styles["SectionHeading"]),
                Paragraph(DISCLAIMER_TEXT, styles["BodySmall"]),
            ]
        )

        document = SimpleDocTemplate(str(target), pagesize=A4, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
        document.build(story)
        return target

    def _add_heading(self, document: Document, text: str) -> None:
        heading = document.add_heading(text, level=1)
        for run in heading.runs:
            run.font.size = Pt(14)


def _set_normal_font(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)


def _executive_summary_lines(company: CompanyDashboardView) -> list[str]:
    return [
        f"Listed status: {company.listed_status}",
        f"Revenue: {_render_number(company.revenue)} crore",
        f"Latest rating: {company.ratings.rating or 'Not available'}",
        f"TReDS status: {company.tred.status}",
        f"Simulation: {company.simulation.status}",
        f"Risk indicator: {company.risk_indicator.label}",
    ]


def _financial_metric_rows(company: CompanyDashboardView) -> dict[str, str]:
    summary = company.financials.summary
    return {
        "Revenue (Cr)": _render_number(summary.revenue_crore),
        "EBITDA Margin %": _render_number(summary.ebitda_margin_pct),
        "PAT Margin %": _render_number(summary.pat_margin_pct),
        "Debt / Equity": _render_number(summary.debt_to_equity),
        "Interest Coverage": _render_number(summary.interest_coverage),
        "Working Capital Days": _render_number(summary.working_capital_days),
        "Receivables Days": _render_number(summary.receivables_days),
        "Inventory Days": _render_number(summary.inventory_days),
        "Market Cap (Cr)": _render_number(summary.market_cap_crore),
        "Borrowings (Cr)": _render_number(summary.total_borrowings_crore),
    }


def _treds_lines(company: CompanyDashboardView) -> list[str]:
    return [
        f"Status: {company.tred.status}",
        f"Platforms: {', '.join(company.tred.platforms) or 'Not detected'}",
        f"Signal strength: {company.tred.signal_strength}",
        f"Estimated limit (Cr): {_render_number(company.tred.estimated_limit_crore)}",
        f"Method: {company.tred.method_used or 'Not available'}",
    ] + [f"Mention: {mention}" for mention in company.tred.raw_mentions[:5]]


def _simulation_lines(company: CompanyDashboardView) -> list[str]:
    return [
        f"Required: {'Yes' if company.simulation.required_flag else 'No'}",
        f"Status: {company.simulation.status}",
        f"Simulated rating: {company.simulation.simulated_rating or 'Not available'}",
        f"Range: {company.simulation.rating_range or 'Not available'}",
        f"Confidence: {company.simulation.confidence_label or 'Not available'}",
        f"Method: {company.simulation.method_used or 'Not available'}",
    ] + [f"Note: {note}" for note in company.simulation.notes]


def _build_portfolio_financial_rows(company: CompanyDashboardView) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for statement_type, statement_rows in company.financials.tables.items():
        for row in statement_rows:
            rows.append(
                {
                    "company_id": company.company_id,
                    "company_name": company.company_name,
                    "statement_type": statement_type,
                    **row,
                }
            )
    if rows:
        return rows
    return [
        {
            "company_id": company.company_id,
            "company_name": company.company_name,
            "statement_type": "summary",
            **company.financials.summary.model_dump(),
        }
    ]


def _build_portfolio_ratings_rows(company: CompanyDashboardView) -> list[dict[str, Any]]:
    if not company.ratings.history:
        return [
            {
                "company_id": company.company_id,
                "company_name": company.company_name,
                "rating_date": company.ratings.rating_date,
                "agency": company.ratings.agency,
                "rating": company.ratings.rating,
                "outlook": company.ratings.outlook,
                "rating_action": company.ratings.rating_action,
                "source_url": None,
                "source": "summary",
            }
        ]

    return [
        {
            "company_id": company.company_id,
            "company_name": company.company_name,
            "rating_date": row.rating_date,
            "agency": row.agency_name,
            "rating": row.rating,
            "outlook": row.outlook,
            "rating_action": row.rating_action,
            "source_url": row.source_url,
            "source": row.source,
        }
        for row in company.ratings.history
    ]


def _safe_name(value: str) -> str:
    return "".join(char if char.isalnum() else "_" for char in value).strip("_") or "company"


def _unique_safe_name(value: str, used_names: set[str], *, fallback: str) -> str:
    base_name = _safe_name(value) or fallback
    candidate = base_name
    suffix = 2
    while candidate in used_names:
        candidate = f"{base_name}__{suffix}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def _render_number(value: Any) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, (int, float)):
        return f"{value:,.2f}"
    return str(value)
