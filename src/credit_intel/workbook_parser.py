from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .schemas import FinancialSummary


SECTION_HEADERS = {
    "PROFIT & LOSS": "profit_loss",
    "Quarters": "quarterly_results",
    "BALANCE SHEET": "balance_sheet",
    "CASH FLOW:": "cash_flow",
}


class ScreenerWorkbookParser:
    def parse(self, workbook_path: Path) -> dict[str, Any]:
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        data_sheet = workbook["Data Sheet"]
        rows = self._load_rows(data_sheet)
        meta = self._parse_meta(rows)
        section_rows = self._split_sections(rows)
        profit_loss = self._parse_statement(section_rows.get("profit_loss", []))
        quarterly = self._parse_statement(section_rows.get("quarterly_results", []))
        balance_sheet = self._parse_statement(section_rows.get("balance_sheet", []))
        cash_flow = self._parse_statement(section_rows.get("cash_flow", []))
        financial_summary = self._build_financial_summary(meta, profit_loss, balance_sheet)
        return {
            "meta": meta,
            "financial_summary": financial_summary.model_dump(),
            "profit_loss": profit_loss,
            "quarterly_results": quarterly,
            "balance_sheet": balance_sheet,
            "cash_flow": cash_flow,
        }

    def _load_rows(self, worksheet) -> list[list[Any]]:  # type: ignore[no-untyped-def]
        rows: list[list[Any]] = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append(list(row))
        return rows

    def _parse_meta(self, rows: list[list[Any]]) -> dict[str, Any]:
        meta: dict[str, Any] = {}
        for row in rows[:14]:
            label = self._string(row[0])
            value = row[1] if len(row) > 1 else None
            if label:
                meta[label] = self._clean_value(value)
        return meta

    def _split_sections(self, rows: list[list[Any]]) -> dict[str, list[list[Any]]]:
        sections: dict[str, list[list[Any]]] = {}
        current_key: str | None = None
        for row in rows:
            label = self._string(row[0])
            if label in SECTION_HEADERS:
                current_key = SECTION_HEADERS[label]
                sections[current_key] = []
                continue
            if current_key:
                sections[current_key].append(row)
        return sections

    def _parse_statement(self, rows: list[list[Any]]) -> list[dict[str, Any]]:
        if not rows:
            return []
        header_row = next((row for row in rows if self._string(row[0]) == "Report Date"), None)
        if not header_row:
            return []

        periods = [self._serialize_date(value) for value in header_row[1:] if value is not None]
        records = [{"period": period} for period in periods]
        for row in rows:
            label = self._string(row[0])
            if not label or label == "Report Date":
                continue
            values = row[1 : len(periods) + 1]
            metric_key = label.strip().lower().replace("&", "and").replace("%", "pct").replace(" ", "_")
            for record, raw_value in zip(records, values, strict=False):
                record[metric_key] = self._clean_value(raw_value)
        return records

    def _build_financial_summary(
        self,
        meta: dict[str, Any],
        profit_loss: list[dict[str, Any]],
        balance_sheet: list[dict[str, Any]],
    ) -> FinancialSummary:
        latest_pnl = profit_loss[-1] if profit_loss else {}
        latest_balance = balance_sheet[-1] if balance_sheet else {}
        sales = _coerce_float(latest_pnl.get("sales"))
        operating_profit = _coerce_float(latest_pnl.get("operating_profit"))
        net_profit = _coerce_float(latest_pnl.get("net_profit"))
        other_income = _coerce_float(latest_pnl.get("other_income"))
        interest = _coerce_float(latest_pnl.get("interest"))
        borrowings = _coerce_float(latest_balance.get("borrowings"))
        equity = _coerce_float(latest_balance.get("equity_share_capital"))
        reserves = _coerce_float(latest_balance.get("reserves"))

        ebitda_margin = round((operating_profit / sales) * 100, 2) if sales and operating_profit is not None else None
        pat_margin = round((net_profit / sales) * 100, 2) if sales and net_profit is not None else None
        debt_to_equity = round(borrowings / (equity + reserves), 2) if borrowings is not None and equity is not None and reserves is not None and (equity + reserves) else None
        interest_coverage = round((operating_profit + (other_income or 0.0)) / interest, 2) if operating_profit is not None and interest not in {None, 0} else None

        return FinancialSummary(
            revenue_crore=sales,
            ebitda_margin_pct=ebitda_margin,
            pat_margin_pct=pat_margin,
            debt_to_equity=debt_to_equity,
            interest_coverage=interest_coverage,
            working_capital_days=None,
            receivables_days=None,
            inventory_days=None,
            current_price=_coerce_float(meta.get("Current Price")),
            market_cap_crore=_coerce_float(meta.get("Market Capitalization")),
            networth_crore=(_coerce_float(latest_balance.get("equity_share_capital")) or 0.0) + (_coerce_float(latest_balance.get("reserves")) or 0.0) if latest_balance else None,
            total_borrowings_crore=borrowings,
            source="screener_workbook",
            statement_period=self._string(latest_pnl.get("period")),
        )

    def _string(self, value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    def _serialize_date(self, value: Any) -> str | None:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return self._string(value)

    def _clean_value(self, value: Any) -> Any:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if value is None:
            return None
        return value


def _coerce_float(value: Any) -> float | None:
    try:
        if value is None:
            return None
        return float(value)
    except (TypeError, ValueError):
        return None
