from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .config import REFERENCE_SHEET_NAMES


def load_reference_sheet_names(reference_path: Path) -> list[str]:
    workbook = load_workbook(reference_path, read_only=True)
    return workbook.sheetnames


def validate_excel_structure(workbook_path: Path) -> tuple[bool, list[str]]:
    workbook = load_workbook(workbook_path, read_only=True)
    sheet_names = workbook.sheetnames
    missing = [sheet for sheet in REFERENCE_SHEET_NAMES if sheet not in sheet_names]
    return (not missing, missing)

